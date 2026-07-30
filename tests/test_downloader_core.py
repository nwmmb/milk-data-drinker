from __future__ import annotations

import threading
from datetime import date
from pathlib import Path

import pytest
import requests
from openpyxl import load_workbook

from timeless_downloader import core
from timeless_downloader.core import (
    DownloadSettings,
    calendar_ranges,
    date_ranges,
    default_window,
    output_filename,
    planned_downloads,
    previous_full_month,
    run_download,
    validate_settings,
    week_ranges,
)


class FakeResponse:
    def __init__(self, status=200, content=b"report", url="https://example.test/report"):
        self.status_code = status
        self.content = content
        self.url = url
        self.history = []
        self.headers = {}

    def raise_for_status(self):
        if self.status_code >= 400:
            raise requests.HTTPError(
                f"{self.status_code} error", response=self
            )


class FakeSession:
    def __init__(self, responses):
        self.responses = list(responses)
        self.calls = []
        self.closed = False

    def get(self, url, params, timeout):
        self.calls.append((url, params, timeout))
        response = self.responses.pop(0)
        if isinstance(response, Exception):
            raise response
        return response

    def close(self):
        self.closed = True


def settings(tmp_path: Path, **changes) -> DownloadSettings:
    values = {
        "report_type": "deposit_record",
        "start_date": date(2025, 1, 1),
        "end_date": date(2025, 1, 1),
        "window_size": "weekly",
        "output_dir": tmp_path / "downloads",
        "cookie": "session=super-secret",
    }
    values.update(changes)
    return DownloadSettings(**values)


def test_previous_full_month_handles_year_boundary():
    assert previous_full_month(date(2026, 1, 4)) == (
        date(2025, 12, 1),
        date(2025, 12, 31),
    )


def test_exact_date_validation(tmp_path):
    validate_settings(settings(tmp_path), today=date(2025, 1, 2))
    with pytest.raises(ValueError, match="start date"):
        validate_settings(
            settings(tmp_path, start_date=date(2025, 1, 2)),
            today=date(2025, 1, 2),
        )
    with pytest.raises(ValueError, match="between"):
        validate_settings(
            settings(tmp_path, end_date=date(2025, 1, 3)),
            today=date(2025, 1, 2),
        )


def test_weekly_ranges_are_seven_day_chunks_with_partial_last():
    assert list(week_ranges(date(2025, 1, 3), date(2025, 1, 12))) == [
        (date(2025, 1, 3), date(2025, 1, 9)),
        (date(2025, 1, 10), date(2025, 1, 12)),
    ]


def test_calendar_ranges_align_to_calendar_quarters():
    assert list(calendar_ranges(date(2025, 2, 15), date(2025, 8, 2), 3)) == [
        (date(2025, 2, 15), date(2025, 3, 31)),
        (date(2025, 4, 1), date(2025, 6, 30)),
        (date(2025, 7, 1), date(2025, 8, 2)),
    ]


def test_report_defaults_and_planned_filter_count(tmp_path):
    assert default_window("wastage_report") == "weekly"
    assert default_window("batch_summary") == "weekly"
    assert default_window("donor_information") == "quarterly"
    donor = settings(
        tmp_path,
        report_type="donor_information",
        start_date=date(2025, 1, 1),
        end_date=date(2025, 3, 31),
        window_size="quarterly",
    )
    assert date_ranges(donor) == [(date(2025, 1, 1), date(2025, 3, 31))]
    assert planned_downloads(donor) == 4


def test_output_filename_uses_exact_dates_and_filter():
    assert output_filename(
        "donor_information",
        {"name": "active"},
        date(2025, 2, 15),
        date(2025, 3, 31),
    ) == "donor_information_active_2025-02-15_to_2025-03-31.xls"


def test_preview_needs_no_cookie_or_output_and_orders_callbacks(tmp_path):
    events = []
    preview = settings(
        tmp_path,
        start_date=date(2025, 1, 1),
        end_date=date(2025, 1, 8),
        preview=True,
        cookie="",
    )

    def should_not_make_session(_cookie):
        raise AssertionError("preview created a session")

    result = run_download(
        preview, callback=events.append, session_factory=should_not_make_session
    )

    assert result.planned == result.completed == 2
    assert not preview.output_dir.exists()
    assert events[0].kind == "started"
    assert events[-1].kind == "complete"
    assert [event.kind for event in events].count("progress") == 2


def test_success_existing_no_data_and_no_cookie_in_logs(tmp_path):
    run_settings = settings(
        tmp_path,
        start_date=date(2025, 1, 1),
        end_date=date(2025, 1, 15),
    )
    run_settings.output_dir.mkdir()
    first = run_settings.output_dir / output_filename(
        "deposit_record", {}, date(2025, 1, 1), date(2025, 1, 7)
    )
    first.write_bytes(b"existing")
    session = FakeSession(
        [
            FakeResponse(content=b"No data found matching report criteria"),
            FakeResponse(content=b"new report"),
        ]
    )
    events = []
    result = run_download(
        run_settings,
        callback=events.append,
        session_factory=lambda _cookie: session,
        delay=0,
    )

    assert (result.existing, result.no_data, result.downloaded, result.failed) == (
        1,
        1,
        1,
        0,
    )
    assert result.completed == result.planned == 3
    assert session.closed
    assert "super-secret" not in "\n".join(event.message for event in events)
    assert events[0].kind == "started"
    assert events[-1].kind == "complete"


def test_server_errors_retry_then_succeed(tmp_path):
    session = FakeSession(
        [FakeResponse(status=500), FakeResponse(status=503), FakeResponse(content=b"ok")]
    )
    result = run_download(
        settings(tmp_path),
        session_factory=lambda _cookie: session,
        delay=0,
        retry_backoff=(0, 0, 0),
    )
    assert result.downloaded == 1
    assert result.failed == 0
    assert len(session.calls) == 3


def test_exhausted_server_errors_are_failed_not_fatal(tmp_path):
    session = FakeSession([FakeResponse(status=500) for _ in range(3)])
    result = run_download(
        settings(tmp_path),
        session_factory=lambda _cookie: session,
        delay=0,
        retry_backoff=(0, 0, 0),
    )
    assert result.failed == 1
    assert result.fatal_error is None


def test_timeout_fallback_uses_daily_requests_inside_parent_progress(tmp_path):
    run_settings = settings(
        tmp_path, start_date=date(2025, 1, 1), end_date=date(2025, 1, 2)
    )
    session = FakeSession(
        [
            requests.Timeout("weekly timeout"),
            FakeResponse(content=b"day one"),
            FakeResponse(content=b"No data found matching report criteria"),
        ]
    )
    events = []
    result = run_download(
        run_settings,
        callback=events.append,
        session_factory=lambda _cookie: session,
        delay=0,
    )
    assert result.planned == result.completed == 1
    assert result.downloaded == 1
    assert result.no_data == 1
    assert len(session.calls) == 3
    assert len([event for event in events if event.kind == "progress"]) == 1


@pytest.mark.parametrize(
    "response, message",
    [
        (FakeResponse(url="https://example.test/login"), "redirected to login"),
        (
            (lambda response: (response.headers.update({"Location": "/login"}), response)[1])(
                FakeResponse(status=302)
            ),
            "redirected to login",
        ),
        (FakeResponse(status=403), "403 Forbidden"),
    ],
)
def test_authentication_and_403_are_fatal(tmp_path, response, message):
    result = run_download(
        settings(tmp_path),
        session_factory=lambda _cookie: FakeSession([response]),
        delay=0,
    )
    assert message in result.fatal_error
    assert result.failed == 0


def test_cancellation_stops_before_next_request_and_keeps_file(tmp_path):
    run_settings = settings(
        tmp_path, start_date=date(2025, 1, 1), end_date=date(2025, 1, 8)
    )
    session = FakeSession([FakeResponse(content=b"first"), FakeResponse(content=b"second")])
    cancelled = threading.Event()

    def callback(event):
        if event.kind == "counts" and event.downloaded == 1:
            cancelled.set()

    result = run_download(
        run_settings,
        callback=callback,
        cancel_event=cancelled,
        session_factory=lambda _cookie: session,
        delay=0,
    )
    assert result.cancelled
    assert result.downloaded == 1
    assert result.completed == 1
    assert len(result.files) == 1
    assert result.files[0].read_bytes() == b"first"
    assert len(session.calls) == 1


def test_combine_runs_only_when_selected(tmp_path, monkeypatch):
    combined = tmp_path / "combined.xlsx"
    calls = []

    def fake_combine(*args, **kwargs):
        calls.append((args, kwargs))
        return combined

    monkeypatch.setattr(core, "combine_files", fake_combine)
    result = run_download(
        settings(tmp_path, combine=True),
        session_factory=lambda _cookie: FakeSession([FakeResponse()]),
        delay=0,
    )
    assert result.combined_file == combined
    assert len(calls) == 1

    calls.clear()
    result = run_download(
        settings(tmp_path, output_dir=tmp_path / "other", combine=False),
        session_factory=lambda _cookie: FakeSession([FakeResponse()]),
        delay=0,
    )
    assert result.combined_file is None
    assert calls == []


def _deposit_source(path, columns=None, deposit_id="DEP000001"):
    columns = columns or [
        "Deposit", "Donor/Milk Bank", "Expiry Date", "Volume Remaining (mL)"
    ]
    rows = [
        {
            "Deposit": deposit_id,
            "Donor/Milk Bank": "DON000001: Synthetic Donor",
            "Expiry Date": "2027-01-01",
            "Volume Remaining (mL)": 100,
        },
        {"Deposit": "Summary"},
        {"Deposit": "Total"},
    ]
    import pandas as pd

    pd.DataFrame(rows, columns=columns).to_csv(path, index=False)


def test_strict_combine_preserves_order_and_writes_hidden_metadata(tmp_path):
    first = tmp_path / "deposit_record_first.xls"
    second = tmp_path / "deposit_record_second.xls"
    _deposit_source(first, deposit_id="DEP000001")
    _deposit_source(second, deposit_id="DEP000002")

    output = core.combine_files(
        "deposit_record",
        [first, second],
        date(2026, 7, 1),
        date(2026, 7, 31),
        tmp_path,
    )

    workbook = load_workbook(output, data_only=True)
    assert workbook.sheetnames[0] == "Combined Data"
    assert workbook["_timeless_downloader_metadata"].sheet_state == "hidden"
    metadata = dict(workbook["_timeless_downloader_metadata"].values)
    assert metadata == {
        "artifact_type": "combined",
        "report_type": "deposit_record",
        "format_version": "1",
    }
    assert [cell.value for cell in workbook["Combined Data"][1]] == [
        "Deposit", "Donor/Milk Bank", "Expiry Date", "Volume Remaining (mL)"
    ]
    workbook.close()


def test_strict_combine_rejects_ordered_column_mismatch(tmp_path):
    first = tmp_path / "first.xls"
    second = tmp_path / "second.xls"
    columns = ["Deposit", "Donor/Milk Bank", "Expiry Date", "Volume Remaining (mL)"]
    _deposit_source(first, columns)
    _deposit_source(second, list(reversed(columns)))
    with pytest.raises(ValueError, match="ordered source columns differ"):
        core.combine_files(
            "deposit_record", [first, second], date(2026, 7, 1), date(2026, 7, 31), tmp_path
        )


def test_strict_combine_rejects_a_different_report_type(tmp_path):
    wrong = tmp_path / "wrong.xls"
    wrong.write_text(
        "Order Number,Recipient,Bottle Size (mL),Dispense Date\n"
        "ORD1,HOS1: Facility,120,2026-07-01\n"
    )
    with pytest.raises(ValueError, match="dispensation"):
        core.combine_files(
            "deposit_record", [wrong], date(2026, 7, 1), date(2026, 7, 31), tmp_path
        )
